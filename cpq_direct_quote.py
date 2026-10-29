"""普通报价内的六步确认。只从持久化步骤定价，不创建技术项目。"""
from __future__ import annotations

import copy
import io
import json
import re
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP, localcontext

KEY = '_direct_quote'
MODE = 'direct_quote_v1'
NAMES = ['确认需求配置', '报价成本确认', '定价-利润加成', '报价-其他加价项', '报价方案', '输出报价单']
PROMPT_OPEN = '【直接报价参数】'
PROMPT_CLOSE = '【/直接报价参数】'


def new_state():
    return {'mode': MODE, 'next_step': 1, 'steps': {}, 'quote': None}


def number(value, label, *, minimum=Decimal(0), maximum=Decimal('1000000000000')):
    if isinstance(value, bool) or value is None or str(value).strip() == '':
        raise ValueError(label + '必须填写有效数字')
    try:
        result = Decimal(str(value))
    except InvalidOperation:
        raise ValueError(label + '不是有效数字') from None
    if not result.is_finite() or result < minimum or result > maximum:
        raise ValueError(label + '超出允许范围')
    if result.as_tuple().exponent < -12 or len(result.as_tuple().digits) > 24:
        raise ValueError(label + '精度过高，最多支持12位小数')
    return result


def text(data, key, label, *, required=True):
    value = data.get(key, '')
    if not isinstance(value, str) or len(value) > 20000 or (required and not value.strip()):
        raise ValueError('请填写' + label)
    return value.strip()


def normalize(step, data):
    if not isinstance(data, dict):
        raise ValueError('步骤输入必须是对象')
    if step == 1:
        quantity = number(data.get('quantity'), '数量', minimum=Decimal(1), maximum=Decimal('1000000000'))
        if quantity != quantity.to_integral_value():
            raise ValueError('数量必须是正整数')
        return {**{k: text(data, k, label) for k, label in
                   [('product_name', '产品名称'), ('customer', '客户'), ('specification', '规格材料')]},
                'product_code': text(data, 'product_code', '产品编码', required=False),
                'quantity': int(quantity)}
    if step == 2:
        cost = number(data.get('unit_cost'), '每件未税成本')
        if cost <= 0:
            raise ValueError('每件未税成本必须大于零')
        if data.get('confirmed') is not True:
            raise ValueError('请核实成本及来源后显式确认')
        if not isinstance(data.get('demo'), bool):
            raise ValueError('请选择实际核实成本或演示成本')
        currency = text(data, 'currency', '币种').upper()
        if not re.fullmatch('[A-Z]{3}', currency):
            raise ValueError('币种需为三位代码，如CNY')
        return {'unit_cost': str(cost), 'currency': currency,
                'cost_source': text(data, 'cost_source', '成本来源'), 'confirmed': True, 'demo': data['demo']}
    if step == 3:
        margin = number(data.get('margin_percent'), '毛利率', maximum=Decimal(100))
        if margin >= 100:
            raise ValueError('毛利率必须小于100%')
        return {'margin_percent': str(margin)}
    if step == 4:
        return {'addon_per_unit': str(number(data.get('addon_per_unit'), '每件未税其他加价')),
                'tax_percent': str(number(data.get('tax_percent'), '税率', maximum=Decimal(100)))}
    return {}  # 不接受客户端价格或报价单正文。


def parse_prompt(prompt):
    if isinstance(prompt, str) and PROMPT_OPEN not in prompt and PROMPT_CLOSE not in prompt:
        from cpq_customer_quote import parse_customer_prompt
        return parse_customer_prompt(prompt)
    if (not isinstance(prompt, str) or len(prompt) > 100000
            or prompt.count(PROMPT_OPEN) != 1 or prompt.count(PROMPT_CLOSE) != 1):
        raise ValueError('请提供唯一且完整的【直接报价参数】JSON参数块')
    start = prompt.index(PROMPT_OPEN) + len(PROMPT_OPEN)
    end = prompt.index(PROMPT_CLOSE)
    if end < start:
        raise ValueError('直接报价参数块顺序错误')

    def pairs(items):
        result = {}
        for key, value in items:
            if key in result:
                raise ValueError('直接报价参数存在重复键：' + key)
            result[key] = value
        return result

    def invalid_constant(value):
        raise ValueError('直接报价参数不能含非有限数值：' + value)

    try:
        data = json.loads(prompt[start:end], object_pairs_hook=pairs, parse_constant=invalid_constant)
    except json.JSONDecodeError:
        raise ValueError('直接报价参数不是合法JSON') from None
    if not isinstance(data, dict) or set(data) != {'1', '2', '3', '4'}:
        raise ValueError('直接报价参数必须包含1需求、2成本、3毛利率、4加价税率四个分区')
    inputs = {str(step): normalize(step, data[str(step)]) for step in range(1, 5)}
    calculate(inputs)  # 在启用前验证整条计算路径，避免留下半启用报价。
    return inputs


def calculate(steps):
    def money(value):
        return value.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    with localcontext() as ctx:
        ctx.prec = 50
        base, cost, profit, addon = (steps[str(i)] for i in range(1, 5))
        unit = money(Decimal(cost['unit_cost']) / (1 - Decimal(profit['margin_percent']) / 100)
                     + Decimal(addon['addon_per_unit']))
        if unit <= 0:
            raise ValueError('计算后的成交单价必须至少为0.01')
        total = money(unit * base['quantity'])
        tax = money(total * Decimal(addon['tax_percent']) / 100)
        return {**base, 'currency': cost['currency'], 'demo': cost['demo'],
                'tax_percent': addon['tax_percent'], 'net_unit_price': str(unit),
                'net_total': str(total), 'tax_amount': str(tax), 'gross_total': str(total + tax)}


def advance(state, step, data):
    if not isinstance(state, dict) or state.get('mode') != MODE:
        raise ValueError('本单尚未启用直接报价')
    if isinstance(step, bool) or not isinstance(step, int) or not 1 <= step <= 6:
        raise ValueError('步骤号无效')
    values = normalize(step, data)
    if state.get('auto_pricing'):
        if step >= 2 and state.get('pricing_gaps'):
            raise ValueError('系统报价依据尚未齐备：'+'；'.join(state['pricing_gaps']))
        expected = normalize(step, state['prompt_inputs'].get(str(step), {}))
        if values != expected:
            raise ValueError('本单采用系统绑定的需求和报价依据；修改规格请在第1步重新发送需求')
    if step < state['next_step']:
        if state['steps'].get(str(step)) == values:
            return copy.deepcopy(state)
        raise ValueError('已确认步骤不可覆盖，请新建报价以修改')
    if step != state['next_step']:
        raise ValueError('请先完成前面的报价步骤')
    result = copy.deepcopy(state)
    result['steps'][str(step)] = values
    if step >= 4:
        result['quote'] = calculate(result['steps'])
    result['next_step'] = step + 1
    return result


def state_for(session_id, conn=None):
    import cpq_wf
    snap = cpq_wf.step_snapshot(session_id, 1, conn=conn) or {}
    state = snap.get(KEY) if isinstance(snap, dict) else None
    return state if isinstance(state, dict) and state.get('mode') == MODE else None


def authorized(card, user):
    return bool(card and user and user.get('role_code') == 'sales_mgr'
                and str(card.get('creator_user_id')) == str(user.get('user_id')))


def _store(conn, session_id, state):
    import cpq_auth
    cpq_auth._exec(conn, "UPDATE cpq_wf_card_step SET data_snapshot = "
                   "jsonb_set(COALESCE(data_snapshot, '{}'::jsonb), '{_direct_quote}', %s::jsonb) "
                   "WHERE card_id = (SELECT card_id FROM cpq_wf_card WHERE session_id = %s) AND step_no = 1",
                   (json.dumps(state, ensure_ascii=False), session_id))


def command(session_id, user, action='read', *, step=None, data=None):
    import cpq_auth
    import cpq_wf
    if not isinstance(session_id, str) or not session_id:
        raise ValueError('缺少报价会话')
    prepared = parse_prompt((data or {}).get('prompt')) if action == 'prepare' and isinstance(data, dict) else None
    if action == 'prepare' and prepared is None:
        raise ValueError('缺少直接报价提示词')
    auto = bool(prepared is not None and '_customer_requirement' in prepared)
    metadata = {}
    if auto:
        from cpq_customer_quote import prepare_customer
        prepared = prepare_customer(prepared)
        metadata = dict(auto_pricing=True, pricing_gaps=prepared.pop('_pricing_gaps'),
                        pricing_evidence=prepared.pop('_pricing_evidence'),
                        customer_requirement=prepared.pop('_customer_requirement'))
    conn = cpq_wf.tx_connect()
    try:
        with conn.transaction():
            # 同一卡片的直接报价命令串行化，状态与步骤在一个事务内提交。
            cpq_auth._exec(conn, 'SELECT card_id FROM cpq_wf_card WHERE session_id = %s FOR UPDATE', (session_id,))
            card = cpq_wf._fetch_card(conn, session_id)
            if not authorized(card, user):
                raise cpq_wf.WfError('仅创建本单的销售经理可操作直接报价')
            state = state_for(session_id, conn)
            if action in ('enable', 'prepare') and state is None:
                active = cpq_auth._exec(conn,
                    "SELECT step_no FROM cpq_wf_card_step WHERE card_id = %s AND status <> 'pending' "
                    "AND NOT (step_no = 1 AND status = 'in_progress')",
                    (int(card['card_id']),)).fetchall()
                tasks = cpq_auth._exec(conn,
                    "SELECT task_id FROM cpq_wf_task WHERE card_id = %s AND status IN ('open', 'claimed')",
                    (int(card['card_id']),)).fetchall()
                if (active or tasks or int(card.get('current_step') or 1) != 1
                        or card.get('overall_status') not in ('draft', 'in_progress')):
                    raise ValueError('已有步骤或未结任务，请新建报价再启用直接报价')
                state = new_state()
                if prepared is not None:
                    state['prompt_inputs'] = prepared
                    state.update(metadata)
                _store(conn, session_id, state)
                cpq_wf._log(conn, int(card['card_id']), None, int(user['user_id']),
                            'direct_quote_enabled', 1, 1, '销售经理启用本单直接报价')
            elif action == 'prepare':
                if state['next_step'] != 1 and (state.get('prompt_inputs') != prepared
                                               or bool(state.get('auto_pricing')) != auto):
                    raise ValueError('已有确认步骤，不能替换提示词参数，请新建报价')
                if state['next_step'] == 1 and (state.get('prompt_inputs') != prepared or auto):
                    state = copy.deepcopy(state)
                    state['prompt_inputs'] = prepared
                    for key in ('auto_pricing','pricing_gaps','pricing_evidence','customer_requirement'):
                        state.pop(key, None)
                    state.update(metadata)
                    _store(conn, session_id, state)
            elif action == 'complete':
                updated = advance(state, step, data)
                if updated != state:
                    cpq_wf.complete_step(session_id, step, user,
                        json.dumps({'direct_quote_step': updated['steps'][str(step)]}, ensure_ascii=False),
                        conn=conn, _direct_validated=True)
                    state = updated
                    _store(conn, session_id, state)
            elif action not in ('read', 'enable'):
                raise ValueError('未知直接报价操作')
            return {'ok': True, 'state': state, 'card': cpq_wf._fetch_card(conn, session_id)}
    finally:
        conn.close()


def document(state):
    if not state or state.get('next_step') != 7:
        raise ValueError('请先确认完成六步报价')
    from docx import Document
    quote = calculate(state['steps'])
    doc = Document()
    doc.add_heading('报价单', 0)
    if quote['demo']:
        doc.add_paragraph('演示报价，不用于正式交易')
    table = doc.add_table(rows=0, cols=2)
    for key, label in [('customer','客户'), ('product_name','产品名称'), ('product_code','产品编码'),
                       ('specification','规格材料'), ('quantity','数量'), ('currency','币种'),
                       ('net_unit_price','未税单价'), ('net_total','未税合计'),
                       ('tax_percent','税率（%）'), ('tax_amount','税额'), ('gross_total','含税合计')]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, str(quote.get(key, ''))
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()
