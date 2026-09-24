"""包装 DWG 的业务部件层与 CAD 平面图查看器红测。"""
from pathlib import Path
import re
import unittest

from openpyxl import load_workbook

from tech_app.backend.services import packaging_parts


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "裕同包装项目-待开发" / "酒盒 报价资料.xlsx"
APP_JS = (ROOT / "tech_app/frontend/app.js").read_text(encoding="utf-8")
INDEX_HTML = (ROOT / "tech_app/frontend/index.html").read_text(encoding="utf-8")
PARTS_PY = (ROOT / "tech_app/backend/services/packaging_parts.py").read_text(encoding="utf-8")
BOM_PY = (ROOT / "tech_app/backend/services/packaging_bom.py").read_text(encoding="utf-8")
COST_PY = (ROOT / "tech_app/backend/services/packaging_cost.py").read_text(encoding="utf-8")


class WineBoxAuthorityWorkbookFacts(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wb = load_workbook(WORKBOOK, data_only=False)
        cls.ws = cls.wb[cls.wb.sheetnames[0]]

    def test_workbook_really_contains_28_business_rows_and_28_part_images(self):
        rows = []
        for row in range(4, 32):
            if isinstance(self.ws.cell(row, 1).value, int):
                rows.append((self.ws.cell(row, 1).value, self.ws.cell(row, 2).value))
        self.assertEqual(28, len(rows))
        self.assertEqual(list(range(1, 29)), [row[0] for row in rows])
        self.assertEqual(28, len(self.ws._images))

    def test_note_and_author_rows_are_not_parts(self):
        # 第 32 行虽然 A 列误填了序号 29，但 B 列是整句客户要求，不是部件名；
        # 导入器不能只凭“序号是整数”就把它吞成第 29 个部件。
        self.assertEqual(29, self.ws.cell(32, 1).value)
        self.assertIn("干燥剂", str(self.ws.cell(32, 2).value))
        self.assertIn("制表", str(self.ws.cell(33, 1).value))
        self.assertIsNone(self.ws.cell(32, 3).value)
        self.assertIsNone(self.ws.cell(33, 3).value)


class BusinessPartLayerRed(unittest.TestCase):
    def test_authoritative_workbook_importer_exists(self):
        module = ROOT / "tech_app/backend/services/packaging_reference_workbook.py"
        self.assertTrue(module.exists(), "缺少包装业务部件权威资料导入器")
        if module.exists():
            source = module.read_text(encoding="utf-8")
            self.assertIn("import_workbook", source)
            self.assertIn("business_part_code", source)
            self.assertIn("thumbnail_ref", source)
            self.assertIn("merged_from", source)

    def test_parts_document_has_business_parts_and_separate_geometry_evidence(self):
        self.assertIn("packaging-business-parts/1", PARTS_PY)
        self.assertIn('"business_parts"', PARTS_PY)
        self.assertIn('"geometry_evidence"', PARTS_PY)
        self.assertIn('"business_part_total"', PARTS_PY)

    def test_geometry_part_code_is_not_the_business_identity(self):
        self.assertNotRegex(PARTS_PY, r'part_code\s*=\s*PART_CODE_FORMAT\s*%\s*index')
        self.assertIn("geometry_component_ref", PARTS_PY)

    def test_missing_authority_does_not_fallback_to_hundreds_of_business_parts(self):
        self.assertIn("business_parts_missing", PARTS_PY)
        self.assertIn("尚未形成业务部件清单", PARTS_PY)

    def test_binding_supports_many_components_per_business_part(self):
        self.assertIn('"component_ids"', PARTS_PY)
        self.assertIn('"entity_ids"', PARTS_PY)
        self.assertIn('"bound_by"', PARTS_PY)
        self.assertIn('"ambiguous"', PARTS_PY)

    def test_downstream_consumes_business_parts_version(self):
        for name, source in (("BOM", BOM_PY), ("成本", COST_PY)):
            self.assertIn("business_parts_id", source, f"{name}没有绑定业务部件版本")
            self.assertIn("business_parts", source, f"{name}仍未切换到业务部件集合")


class PackagingCadPlanViewerRed(unittest.TestCase):
    def test_packaging_mode_has_dedicated_cad_plan_viewer(self):
        combined = INDEX_HTML + APP_JS
        self.assertIn("packagingCadPlanViewer", combined)
        self.assertIn("CAD 平面图", combined)
        self.assertIn("fitPackagingCadPlan", combined)
        self.assertIn("highlightPackagingBusinessPart", combined)

    def test_packaging_label_is_not_hard_coded_as_3d(self):
        self.assertNotIn('<section class="drawing-model-column" aria-label="3D 视图">', INDEX_HTML)
        self.assertNotIn('id="viewerPartName" class="center-title">3D 视图', INDEX_HTML)
        self.assertRegex(APP_JS, r"packaging[\s\S]{0,500}(CAD 平面图|包装展开图)")

    def test_packaging_actions_do_not_offer_flat_extrusion_as_primary_view(self):
        actions_start = APP_JS.index("function packagingPartActionsHtml")
        actions_end = APP_JS.index("function packagingPartSolidReason", actions_start)
        actions = APP_JS[actions_start:actions_end]
        self.assertNotIn("3D 预览", actions)
        self.assertNotIn("packagingPartSolid", actions)

    def test_packaging_selection_highlights_original_cad_entities(self):
        self.assertRegex(APP_JS, r"highlightPackagingBusinessPart\s*\([^)]*entity_ids")
        self.assertIn("geometry_binding", APP_JS)
        self.assertIn("component_ids", APP_JS)

    def test_non_packaging_3d_viewer_remains_available(self):
        self.assertIn("function loadSTL", APP_JS)
        self.assertRegex(APP_JS, r"(industry|currentIndustry)[\s\S]{0,500}packaging")


class NoHardcodedWinePartCatalogRed(unittest.TestCase):
    def test_production_source_does_not_hardcode_all_wine_part_names(self):
        # 只检查足够专有的部件名；“磁铁/EVA”等可能合法存在于通用材料分类词表。
        names = ["左盖面纸", "右盖面纸", "左盖外盒里层灰板1", "内盒1灰板", "底托灰板"]
        hits = sum(name in PARTS_PY for name in names)
        self.assertEqual(0, hits, "业务部件名称必须由权威资料导入，不能硬编码进拆件引擎")


if __name__ == "__main__":
    unittest.main()
